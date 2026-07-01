"""Styled Excel (.xlsx) renderer (openpyxl) for a :class:`ReportDocument`.

Money columns stay numeric (with a thousands number format) so the sheet remains
fully computable; the header row is frozen and styled, and a bold totals row is
appended when the document carries totals.
"""
import re
from io import BytesIO

from app.aman.services.export.report_document import ReportDocument

_MONEY_FMT = "#,##0.00;(#,##0.00)"


def _num(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return value


def render_excel(doc: ReportDocument) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    # Excel forbids \ / ? * [ ] : in sheet titles and caps length at 31.
    ws.title = re.sub(r"[\\/?*\[\]:]", "-", (doc.title or "Report"))[:31]

    ncols = len(doc.columns)
    last_col = get_column_letter(ncols)
    bold = Font(bold=True)
    white_bold = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E293B")
    totals_fill = PatternFill("solid", fgColor="F1F5F9")
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    r = 1
    if doc.company:
        ws.merge_cells(f"A{r}:{last_col}{r}")
        cell = ws.cell(r, 1, doc.company)
        cell.font = Font(bold=True, size=13)
        cell.alignment = Alignment(horizontal="center")
        r += 1
    ws.merge_cells(f"A{r}:{last_col}{r}")
    cell = ws.cell(r, 1, doc.title)
    cell.font = Font(bold=True, size=11, color="1D4ED8")
    cell.alignment = Alignment(horizontal="center")
    r += 1
    if doc.period:
        ws.merge_cells(f"A{r}:{last_col}{r}")
        cell = ws.cell(r, 1, doc.period)
        cell.font = Font(size=9, color="64748B")
        cell.alignment = Alignment(horizontal="center")
        r += 1
    r += 1  # blank spacer row

    # Header row
    header_row = r
    for ci, col in enumerate(doc.columns, start=1):
        cell = ws.cell(header_row, ci, col.label)
        cell.font = white_bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal={"right": "right", "center": "center"}.get(col.align, "left"))
        cell.border = border
    r += 1

    # Data rows
    for row in doc.rows:
        for ci, col in enumerate(doc.columns, start=1):
            raw = row.get(col.key)
            value = _num(raw) if col.fmt == "money" else ("" if raw is None else raw)
            cell = ws.cell(r, ci, value)
            cell.alignment = Alignment(horizontal={"right": "right", "center": "center"}.get(col.align, "left"))
            cell.border = border
            if col.fmt == "money":
                cell.number_format = _MONEY_FMT
        r += 1

    # Totals row
    if doc.totals:
        label_key = doc.totals_label_key or doc.columns[0].key
        for ci, col in enumerate(doc.columns, start=1):
            if col.key == label_key:
                value = doc.totals_label
            elif col.key in doc.totals:
                value = _num(doc.totals[col.key]) if col.fmt == "money" else doc.totals[col.key]
            else:
                value = ""
            cell = ws.cell(r, ci, value)
            cell.font = bold
            cell.fill = totals_fill
            cell.alignment = Alignment(horizontal={"right": "right", "center": "center"}.get(col.align, "left"))
            cell.border = border
            if col.fmt == "money":
                cell.number_format = _MONEY_FMT
        r += 1

    ws.freeze_panes = f"A{header_row + 1}"

    # Column widths from content length
    for ci, col in enumerate(doc.columns, start=1):
        width = len(col.label)
        for row in doc.rows:
            v = row.get(col.key)
            width = max(width, len(str(v)) if v is not None else 0)
        ws.column_dimensions[get_column_letter(ci)].width = min(max(width + 3, 10), 48)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
